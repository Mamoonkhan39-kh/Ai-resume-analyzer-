"""Resume ATS Analyzer — Streamlit + Google Gemini Flash."""
from __future__ import annotations

import io
import json
import os
import re
from typing import Any

import streamlit as st
from docx import Document
from google import genai
from pypdf import PdfReader

APP_TITLE = "Resume ATS Analyzer"
DEFAULT_MODEL = "gemini-3.5-flash-lite"
MAX_FILE_BYTES = 5 * 1024 * 1024


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    """Extract readable text from PDF, DOCX, or TXT uploads."""
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if suffix == "pdf":
            reader = PdfReader(io.BytesIO(file_bytes))
            if getattr(reader, "is_encrypted", False):
                try:
                    if reader.decrypt("") == 0:
                        raise ValueError("This PDF is password-protected. Please upload an unlocked PDF.")
                except Exception as exc:
                    raise ValueError("This PDF is password-protected. Please upload an unlocked PDF.") from exc
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == "docx":
            doc = Document(io.BytesIO(file_bytes))
            chunks = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    chunks.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(chunks)
        elif suffix == "txt":
            text = file_bytes.decode("utf-8-sig", errors="replace")
        else:
            raise ValueError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Could not read this file. Please check that it is a valid {suffix.upper()} file.") from exc

    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 80:
        raise ValueError(
            "Very little text could be extracted. If this is a scanned/image-only PDF, run OCR first, "
            "then upload a text-based PDF or DOCX."
        )
    return text[:30000]


def heuristic_score(resume_text: str, job_description: str = "") -> tuple[int, dict[str, int]]:
    """A transparent baseline checklist score; not a real ATS vendor score."""
    text = resume_text.lower()
    checks = {
        "Contact details": 15 if re.search(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", text) else 0,
        "Phone number": 10 if re.search(r"(?:\+?\d[\d ().-]{7,}\d)", resume_text) else 0,
        "Clear section headings": 15 if sum(h in text for h in ["experience", "education", "skills", "projects"]) >= 3 else 5 if any(h in text for h in ["experience", "education", "skills", "projects"]) else 0,
        "Skills section": 15 if "skills" in text else 0,
        "Experience / projects": 15 if any(k in text for k in ["experience", "projects", "employment", "internship"]) else 0,
        "Measurable achievements": 15 if re.search(r"\b\d+\s*%|\$\s*\d+|\b\d+\s+(users|clients|projects|hours|days|weeks|months|years)\b", text) else 0,
        "Readable length": 15 if 250 <= len(resume_text.split()) <= 1200 else 8 if len(resume_text.split()) >= 120 else 0,
    }
    if job_description.strip():
        words = lambda s: set(re.findall(r"[a-zA-Z][a-zA-Z+#.-]{2,}", s.lower()))
        jd_words = words(job_description)
        resume_words = words(resume_text)
        keywords = {w for w in jd_words if w not in {"with", "and", "the", "for", "you", "your", "our", "are", "will", "this", "that", "from", "have", "has", "work", "role", "job", "team"}}
        overlap = len(keywords & resume_words) / max(1, len(keywords))
        checks["Job-description keyword overlap"] = round(min(15, overlap * 15))
        # Normalize the baseline to 100 when job matching is included.
        raw = sum(checks.values())
        score = round(raw / 115 * 100)
    else:
        score = sum(checks.values())
    return max(0, min(100, score)), checks


def analyze_with_gemini(api_key: str, resume_text: str, job_description: str, model_name: str) -> dict[str, Any]:
    """Ask Gemini for structured, actionable feedback."""
    client = genai.Client(api_key=api_key)
    prompt = f"""You are a careful resume coach. Review the resume text and provide practical ATS-readability feedback.
Do not claim to know the score of a specific ATS vendor. Do not invent facts, skills, employers, dates, or metrics.
If a job description is supplied, compare the resume to it and identify only genuinely relevant missing terms.
Return ONLY valid JSON with this schema:
{{
  "estimated_ats_score": integer 0-100,
  "summary": "2-4 sentence summary",
  "strengths": ["..."],
  "improvements": [{{"priority":"High|Medium|Low","issue":"...","recommendation":"..."}}],
  "keyword_suggestions": ["..."],
  "formatting_tips": ["..."],
  "rewritten_bullet_examples": [{{"before_or_issue":"...","suggested_version":"..."}}]
}}
The score must be an estimate of general ATS readiness, not a guarantee of interview success. Do not reward a keyword merely for appearing if it is not supported by the resume.

JOB DESCRIPTION (may be empty):
{job_description[:10000] or '[Not provided]'}

RESUME TEXT:
{resume_text[:30000]}
"""
    response = client.models.generate_content(
        model=model_name,
        contents=prompt,
        config={"response_mime_type": "application/json", "temperature": 0.2},
    )
    if not response.text:
        raise RuntimeError("Gemini returned an empty response. Please try again.")
    try:
        result = json.loads(response.text)
    except json.JSONDecodeError as exc:
        raise RuntimeError("Gemini returned an unexpected response format. Please try again.") from exc
    if not isinstance(result, dict):
        raise RuntimeError("Gemini returned an unexpected response format. Please try again.")
    score = result.get("estimated_ats_score")
    if not isinstance(score, (int, float)):
        result["estimated_ats_score"] = None
    else:
        result["estimated_ats_score"] = max(0, min(100, round(score)))
    return result


def get_api_key() -> str:
    """Read Streamlit secret first, then local environment variable."""
    try:
        secret = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        secret = ""
    return str(secret or os.getenv("GEMINI_API_KEY", "")).strip()


def main() -> None:
    st.set_page_config(page_title=APP_TITLE, page_icon="📄", layout="wide")
    st.title("📄 Resume ATS Analyzer")
    st.write("Upload a resume to get an estimated ATS-readiness score, practical improvements, and optional job-description matching.")
    st.caption("This is an AI-assisted estimate, not an official score from any applicant-tracking system or a guarantee of hiring outcomes.")

    with st.sidebar:
        st.header("Settings")
        model_name = st.text_input("Gemini model", value=DEFAULT_MODEL, help="Uses the Google Gen AI SDK. Change this only to a model enabled for your API key.").strip() or DEFAULT_MODEL
        st.info("Your resume text is sent to Google Gemini for analysis. Avoid uploading information you do not want processed by a third-party AI service.")

    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"], help="Maximum file size: 5 MB. Scanned PDFs may need OCR first.")
    job_description = st.text_area("Target job description (optional)", height=180, placeholder="Paste the job posting here to get more relevant keyword and skills feedback...")

    if uploaded is not None:
        file_bytes = uploaded.getvalue()
        if len(file_bytes) > MAX_FILE_BYTES:
            st.error("File is larger than 5 MB. Please upload a smaller file.")
            st.stop()
        st.success(f"File selected: {uploaded.name} ({len(file_bytes) / 1024:.0f} KB)")

    analyze_clicked = st.button("Analyze resume", type="primary", disabled=uploaded is None, use_container_width=True)
    if analyze_clicked and uploaded is not None:
        api_key = get_api_key()
        if not api_key:
            st.error("Gemini API key not found. Set GEMINI_API_KEY in your environment or Streamlit Secrets. See README.md.")
            st.stop()
        try:
            with st.spinner("Reading your resume..."):
                resume_text = extract_resume_text(uploaded.name, uploaded.getvalue())
            with st.expander("Preview extracted resume text"):
                st.text(resume_text[:5000] + ("\n… [preview shortened]" if len(resume_text) > 5000 else ""))
            with st.spinner("Gemini is reviewing your resume. This may take a little while..."):
                result = analyze_with_gemini(api_key, resume_text, job_description, model_name)
            baseline, checks = heuristic_score(resume_text, job_description)
            ai_score = result.get("estimated_ats_score")
            st.subheader("Your results")
            c1, c2 = st.columns(2)
            with c1:
                st.metric("Gemini estimated ATS readiness", f"{ai_score}/100" if ai_score is not None else "Not available")
            with c2:
                st.metric("Basic checklist score", f"{baseline}/100")
            st.caption("The checklist is a simple local heuristic. The Gemini estimate is subjective; neither reproduces a proprietary ATS algorithm.")
            if result.get("summary"):
                st.markdown("### Summary")
                st.write(result["summary"])
            for section, title in [("strengths", "What is working well"), ("improvements", "Recommended improvements"), ("keyword_suggestions", "Keyword suggestions"), ("formatting_tips", "Formatting tips"), ("rewritten_bullet_examples", "Example bullet rewrites")]:
                items = result.get(section, [])
                if not isinstance(items, list) or not items:
                    continue
                st.markdown(f"### {title}")
                if section == "improvements":
                    for item in items:
                        if isinstance(item, dict):
                            priority = str(item.get("priority", "Medium"))
                            st.markdown(f"**{priority}: {item.get('issue', 'Improvement')}**")
                            st.write(item.get("recommendation", ""))
                elif section == "rewritten_bullet_examples":
                    for item in items:
                        if isinstance(item, dict):
                            if item.get("before_or_issue"):
                                st.markdown(f"**Current issue:** {item['before_or_issue']}")
                            if item.get("suggested_version"):
                                st.markdown(f"**Suggested version:** {item['suggested_version']}")
                else:
                    for item in items:
                        st.markdown(f"- {item}")
            with st.expander("View basic checklist details"):
                for name, points in checks.items():
                    st.write(f"**{name}:** {points} points")
            report = {"gemini_analysis": result, "basic_checklist_score": baseline, "checklist_details": checks}
            st.download_button("Download analysis as JSON", data=json.dumps(report, indent=2, ensure_ascii=False), file_name="resume_analysis.json", mime="application/json", use_container_width=True)
        except ValueError as exc:
            st.error(str(exc))
        except Exception as exc:
            # Keep technical details out of the public UI; details are available in Streamlit logs.
            st.error("Analysis failed. Check your API key, model access, internet connection, and Streamlit logs, then try again.")
            print(f"Resume analysis error: {type(exc).__name__}: {exc}")

    st.divider()
    st.caption("Privacy reminder: this demo does not save resumes to a database, but uploaded content is processed in memory and sent to Gemini when you click Analyze.")


if __name__ == "__main__":
    main()
